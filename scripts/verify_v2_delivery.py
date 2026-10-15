"""Verify preserved V1 records, V2 demonstrations, browser and export evidence.

Read-only database checks; writes only a compact verification artifact.
"""
import json
import sqlite3
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def readonly(path):
    return sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True)


def main():
    backup = ROOT / "data/backups/before-v2.sqlite3"
    current = ROOT / "data/simulator.sqlite3"
    with readonly(backup) as before, readonly(current) as after:
        for table, keys in [("entities", ["kind", "id"]), ("responses", ["id"]), ("answers", ["response_id", "question_id"])]:
            columns = [row[1] for row in before.execute(f"PRAGMA table_info({table})")]
            old = [dict(zip(columns, row)) for row in before.execute(f"SELECT * FROM {table}")]
            for row in old:
                match = " AND ".join(f"{key}=?" for key in keys)
                saved = after.execute(f"SELECT {','.join(columns)} FROM {table} WHERE {match}", tuple(row[key] for key in keys)).fetchone()
                assert saved == tuple(row[c] for c in columns), f"Original {table} record changed"
        counts = {"preserved_v1_personas": before.execute("SELECT count(*) FROM entities WHERE kind='persona'").fetchone()[0],
                  "preserved_v1_responses": before.execute("SELECT count(*) FROM responses").fetchone()[0],
                  "current_personas": after.execute("SELECT count(*) FROM entities WHERE kind='persona'").fetchone()[0],
                  "current_responses": after.execute("SELECT count(*) FROM responses").fetchone()[0],
                  "schema_version": after.execute("PRAGMA user_version").fetchone()[0]}
        assert counts["schema_version"] == 3
        manifest = json.loads((ROOT / "outputs/v2-demo/manifest.json").read_text(encoding="utf-8"))
        for task_id in manifest["task_ids"]:
            assert after.execute("SELECT status FROM tasks WHERE id=?", (task_id,)).fetchone()[0] == "completed"
            assert after.execute("SELECT count(*) FROM responses WHERE task_id=?", (task_id,)).fetchone()[0] == 6

    exports = {}
    for name in ["model-comparison", "scenario-comparison", "quality", "motivations"]:
        folder = ROOT / "outputs/v2-demo"
        report = json.loads((folder / (name + ".json")).read_text(encoding="utf-8"))
        assert report["dataset"]["dataset_kind"] == "simulation"
        if "comparison" in name:
            assert report["task_settings"]
        with zipfile.ZipFile(folder / (name + ".zip")) as archive:
            assert archive.testzip() is None
            assert json.loads(archive.read("report.json")) == report
        book = load_workbook(folder / (name + ".xlsx"))
        assert "報告來源與方法" in book.sheetnames
        assert all(cell.data_type != "f" for ws in book for row in ws for cell in row)
        exports[name] = {"JSON": "passed", "CSV ZIP": "passed", "Excel": "passed"}
    browser = json.loads((ROOT / "outputs/ui-v2/verification.json").read_text(encoding="utf-8"))
    assert len(browser["pages"]) == 14 and not browser["javascript_errors"]
    download = load_workbook(ROOT / "outputs/ui-v2/motivation-browser.xlsx")
    assert "segments" in download.sheetnames and "報告來源與方法" in download.sheetnames
    result = {"database": counts, "original_entities_responses_answers": "unchanged",
              "demo": {"personas": 6, "tasks": 5, "responses": 30}, "exports": exports,
              "browser_pages": len(browser["pages"]), "browser_excel_download": "passed",
              "external_services": "not live verified", "github_updated": False}
    (ROOT / "outputs/v2-delivery-verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
