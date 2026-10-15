from __future__ import annotations

import io
import json
import zipfile
from collections import Counter

import pandas as pd

from .security import redact
from .personas import similarity


def canonical(value):
    """Order-insensitive form of an answer value, for comparing answers."""
    if isinstance(value, list):
        return sorted(value)
    if isinstance(value, dict):
        return {k: canonical(v) for k, v in value.items()}
    return value


def display_answer(q, value, other_text=""):
    if value is None:
        return ""
    options = {o.id: o.label for o in q.options}
    if other_text:
        options.update({o.id: f"{o.label}：{other_text}" for o in q.options if o.is_other})
    if isinstance(value, list):
        return "、".join(options.get(v, str(v)) for v in value)
    if isinstance(value, dict):
        rows = {r.id: r.label for r in q.rows}
        return "；".join(f"{rows.get(k,k)}：{display_answer(q,v)}" for k, v in value.items())
    return options.get(str(value), str(value))


def safe_cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        value = redact(value)
        if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
            return "'" + value
    return value


def tables(db, task_id=None):
    responses = db.responses(task_id)
    if task_id:
        from .models import TaskSpec
        personas = TaskSpec.model_validate_json(db.task(task_id)["spec"]).personas
    else:
        by_version = {(p.id, p.revision): p for p in db.list("persona")}
        for r in responses:
            by_version.setdefault((r.persona_id, r.persona_snapshot.revision), r.persona_snapshot)
        personas = list(by_version.values())
    ptable = []
    for p in personas:
        ptable.append({"persona_id": p.id, "name": p.name, "archetype": p.archetype, "revision": p.revision,
            **p.background.model_dump(), **p.traits.model_dump(), **{f"behavior_{k}": v for k, v in p.behavior.model_dump().items()},
            "topic_attributes": p.topic_attributes, "core_values": p.core_values, "narrative": p.narrative, "seed": p.seed})
    wide, details, validation, questions = [], [], [], []
    schema_seen = set()
    for r in responses:
        qm = {q.id: q for q in r.questionnaire_snapshot.questions}
        row = {"response_id": r.id, "task_id": r.task_id, "persona_id": r.persona_id, "questionnaire_id": r.questionnaire_id,
               "questionnaire_revision": r.questionnaire_snapshot.revision, "persona_revision": r.persona_snapshot.revision,
               "provider": r.provider_snapshot.provider, "model": r.provider_snapshot.model, "mode": r.mode, "status": r.status}
        row.update(scenario_id=r.scenario.id if r.scenario else None, scenario_name=r.scenario.name if r.scenario else None,
                   elapsed_seconds=r.elapsed_seconds, supersedes_id=r.supersedes_id)
        row.update({a.question_id: a.value for a in r.answers})
        row.update({f"{a.question_id}__other_text": a.other_text for a in r.answers if any(o.is_other for o in qm[a.question_id].options)})
        wide.append(row)
        for a in r.answers:
            details.append({"response_id": r.id, "task_id": r.task_id, "persona_id": r.persona_id, "question_id": a.question_id,
                "title": qm[a.question_id].title, "value": a.value, "other_text": a.other_text, "answer_text": display_answer(qm[a.question_id], a.value, a.other_text),
                "rationale": a.rationale, "uncertain": a.uncertain, "valid": a.valid, "validation": a.validation})
            schema_key = (r.questionnaire_id, r.questionnaire_snapshot.revision, a.question_id)
            if schema_key not in schema_seen:
                schema_seen.add(schema_key)
                questions.append({"questionnaire_id": r.questionnaire_id, "revision": r.questionnaire_snapshot.revision,
                                  **qm[a.question_id].model_dump(mode="json")})
        for issue in r.issues:
            validation.append({"response_id": r.id, "persona_id": r.persona_id, "task_id": r.task_id, **issue.model_dump()})
        for repair in r.repair_history:
            validation.append({"response_id": r.id, "persona_id": r.persona_id, "task_id": r.task_id,
                "question_id": repair["page_id"], "level": "repair", "code": "repair", "message": repair["issue"], "resolution": repair["resolution"]})
    tasks = db.rows("SELECT id,status,created_at,updated_at,spec FROM tasks" + (" WHERE id=?" if task_id else ""), (task_id,) if task_id else ())
    executions = []
    for task in tasks:
        spec = json.loads(task["spec"])
        for item in db.rows("SELECT * FROM task_items WHERE task_id=?", (task["id"],)):
            executions.append({"task_id": task["id"], "persona_id": item["persona_id"], "status": item["status"],
                "attempts": item["attempts"], "error": item["error"], "provider": spec["provider"]["provider"],
                "model": spec["provider"]["model"], "mode": spec["mode"], "seed": spec["seed"],
                "started_at": task["created_at"], "updated_at": task["updated_at"]})
    metadata = {**db.metadata(), "scope_task_id": task_id or "all", "responses": len(responses)}
    output = {"資料集來源": pd.DataFrame([metadata]), "人物": pd.DataFrame(ptable), "回答": pd.DataFrame(wide),
        "回答明細": pd.DataFrame(details), "執行紀錄": pd.DataFrame(executions), "驗證報告": pd.DataFrame(validation,
             columns=["response_id", "persona_id", "task_id", "question_id", "level", "code", "message", "resolution"]),
        "題目對照": pd.DataFrame(questions), "API用量": pd.DataFrame(db.rows("SELECT * FROM usage" + (" WHERE task_id=?" if task_id else ""), (task_id,) if task_id else ())),
        "提交紀錄": pd.DataFrame(db.rows("SELECT s.* FROM submissions s JOIN responses r ON s.response_id=r.id WHERE r.task_id=?", (task_id,)) if task_id else db.rows("SELECT * FROM submissions"))}
    output["實驗設定"] = pd.DataFrame([exp.model_dump(mode="json") for exp in db.list("experiment")
        if task_id is None or any(cell.task_id == task_id for cell in exp.cells)])
    if task_id:
        runs = {r.website_run.id:r.website_run for r in responses if r.website_run}
        for item in db.rows("SELECT checkpoint FROM task_items WHERE task_id=?", (task_id,)):
            saved = json.loads(item["checkpoint"]).get("website_run")
            if saved:
                from .models import WebsiteRun
                runs[saved["id"]] = WebsiteRun.model_validate(saved)
        website_runs = list(runs.values())
    else:
        website_runs = db.list("website_run")
    output["網站操作紀錄"] = pd.DataFrame([r.model_dump(mode="json") for r in website_runs])
    return {name: df.map(safe_cell) if not df.empty else df for name, df in output.items()}


def export_csv_zip(db, task_id=None):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("dataset_metadata.json", json.dumps(db.metadata(), ensure_ascii=False, indent=2))
        for name, df in tables(db, task_id).items():
            z.writestr(f"{name}.csv", df.to_csv(index=False).encode("utf-8-sig"))
    return buffer.getvalue()


def export_excel(db, task_id=None):
    return frames_excel(tables(db, task_id))


def frames_excel(frames):
    # This is application export, not artifact authoring; Pandas is explicitly requested.
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        from openpyxl.styles import Font, PatternFill, Alignment
        for name, df in frames.items():
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            if ws.max_column and not df.empty:
                ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.fill, c.font = PatternFill("solid", fgColor="2443A0"), Font(color="FFFFFF", bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(45, max(14, max(len(str(c.value or "")) for c in list(col)[:80]) + 2))
                for c in col[1:]:
                    c.alignment = Alignment(vertical="top", wrap_text=True)
    return buffer.getvalue()


def export_json(db, task_id=None):
    responses = db.responses(task_id)
    from .models import TaskSpec
    spec = TaskSpec.model_validate_json(db.task(task_id)["spec"]) if task_id else None
    output = {"dataset": db.metadata(), "task_id": task_id, "personas": [p.model_dump(mode="json") for p in (spec.personas if spec else db.list("persona"))],
        "questionnaires": [q.model_dump(mode="json") for q in ([spec.questionnaire] if spec else db.list("questionnaire"))],
        "website_workflows": [w.model_dump(mode="json") for w in ([spec.website_workflow] if spec and spec.website_workflow else [] if spec else db.list("website_workflow"))],
        "responses": [r.model_dump(mode="json") for r in responses],
        "tables": {name: json.loads(df.to_json(orient="records", force_ascii=False)) for name, df in tables(db, task_id).items()}}
    return redact(json.dumps(output, ensure_ascii=False, indent=2)).encode("utf-8")


def distribution(responses, question_id):
    entries = []
    for r in responses:
        q = next((q for q in r.questionnaire_snapshot.questions if q.id == question_id), None)
        a = next((a for a in r.answers if a.question_id == question_id), None)
        if not q or not a or a.validation in {"skipped", "error"} or a.value is None:
            continue
        vals = a.value if isinstance(a.value, list) else [a.value]
        for val in vals:
            if isinstance(val, dict):
                for row, v in val.items():
                    for item in v if isinstance(v, list) else [v]:
                        entries.append({"選項": display_answer(q, item), "人物類型": r.persona_snapshot.archetype, "列": next(x.label for x in q.rows if x.id == row)})
            else:
                entries.append({"選項": display_answer(q, val), "人物類型": r.persona_snapshot.archetype, "列": "整題"})
    df = pd.DataFrame(entries, columns=["選項", "人物類型", "列"])
    return df.groupby(["選項", "人物類型", "列"], dropna=False).size().reset_index(name="回答次數")


def concentration(responses):
    combinations = Counter(json.dumps([(a.question_id, canonical(a.value), a.other_text) for a in r.answers], ensure_ascii=False, sort_keys=True) for r in responses)
    return {"不同答案組合": len(combinations), "最大重複組合人數": max(combinations.values(), default=0),
            "不同組合比例": len(combinations) / len(responses) if responses else 0}


def similarity_matrix(personas):
    return [[similarity(a, b) if a.id != b.id else 1 for b in personas] for a in personas]
