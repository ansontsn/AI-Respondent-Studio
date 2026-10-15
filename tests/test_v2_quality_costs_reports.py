import io
import json
import sqlite3

import httpx
import pytest
from openpyxl import load_workbook

from survey_sim.batch import BatchRunner
from survey_sim.costs import request_estimate, task_estimate, BudgetedProvider, BudgetExceeded
from survey_sim.engine import AnswerEngine
from survey_sim.errors import ProviderError
from survey_sim.migrations import migrate
from survey_sim.models import ProviderConfig, TaskSpec, uid, Answer
from survey_sim.mock_ai import MockProvider
from survey_sim.motivations import motivation_report
from survey_sim.providers import HTTPProvider, usage_recorder, make_provider
from survey_sim.quality import quality_report, regenerate
from survey_sim.reports import report_exports
from survey_sim.security import SecretStore
from survey_sim.storage import Database, dump


def test_budget_blocks_before_request_and_can_resume(db,spec,config,monkeypatch):
    spec.provider = ProviderConfig(provider="openai",model="test",input_price_per_million=1,output_price_per_million=2,rate_source="verified-test",request_interval=0)
    spec.budget_limit = 0
    calls = []
    def factory(c,key,db,task_id):
        def handle(request):
            calls.append(request)
            payload = json.loads(json.loads(request.content)["messages"][1]["content"])
            answers = MockProvider(config).generate("answer_page",payload,"").data
            return httpx.Response(200,json={"choices":[{"message":{"content":json.dumps(answers)}}],"usage":{"prompt_tokens":10,"completion_tokens":20}})
        return HTTPProvider(c,"test-key",usage_recorder(db,task_id),httpx.MockTransport(handle))
    monkeypatch.setattr("survey_sim.batch.make_provider",factory)
    task = db.create_task(spec)
    runner = BatchRunner(db,SecretStore())
    runner.start(task); runner.thread.join(timeout=10)
    assert not calls and db.task(task)["status"]=="budget_paused"
    assert db.rows("SELECT status FROM task_items WHERE task_id=?",(task,))[0]["status"]=="pending"
    db.update_budget(task,10.)
    runner.start(task);runner.thread.join(timeout=10)
    assert db.task(task)["status"]=="completed" and db.responses(task)
    assert db.rows("SELECT * FROM usage WHERE task_id=?",(task,))


def test_unknown_usage_stops_budget_before_next_page(db,spec,config,monkeypatch):
    spec.provider = ProviderConfig(provider="openai",model="test",input_price_per_million=1,output_price_per_million=2,rate_source="verified-test",request_interval=0)
    spec.budget_limit = 10
    calls = []
    def factory(c,key,db,task):
        def handle(request):
            calls.append(request)
            payload = json.loads(json.loads(request.content)["messages"][1]["content"])
            data = MockProvider(config).generate("answer_page",payload,"").data
            return httpx.Response(200,json={"choices":[{"message":{"content":json.dumps(data)}}]})
        return HTTPProvider(c,"test-key",usage_recorder(db,task),httpx.MockTransport(handle))
    monkeypatch.setattr("survey_sim.batch.make_provider",factory)
    task=db.create_task(spec); runner=BatchRunner(db,SecretStore())
    runner.start(task);runner.thread.join(timeout=10)
    assert len(calls)==1 and db.task(task)["status"]=="budget_paused"
    checkpoint=json.loads(db.rows("SELECT checkpoint FROM task_items WHERE task_id=?",(task,))[0]["checkpoint"])
    assert checkpoint["answers"] and checkpoint["next_page"] != "P1"


def test_cost_unknown_is_not_zero_and_mock_zero(spec):
    unknown=ProviderConfig(provider="openai",model="x")
    assert request_estimate(unknown,{"x":1})["high"] is None
    assert task_estimate(spec)["high"]==0
    paid=unknown.model_copy(update={"input_price_per_million":1.,"output_price_per_million":2.,"rate_source":"test"})
    assert request_estimate(paid,{"x":"測試"})["high"]>0


def test_authentication_failure_is_not_format_repair(spec,persona,config):
    class InvalidKey(MockProvider):
        calls=0
        def generate(self,*args):
            self.calls+=1
            raise ProviderError("API Key 無效")
    provider=InvalidKey(config)
    with pytest.raises(ProviderError): AnswerEngine(provider).run("invalid-key",spec,persona)
    assert provider.calls==1


def test_repeated_patterns_review_and_regeneration_preserves_original(db,spec,persona,config):
    spec.mode = "custom"
    spec.variability = .3
    spec.budget_limit = 1.
    people=[]
    for i in range(3):
        p=persona.model_copy(deep=True);p.id=uid();p.name=f"人物{i}";people.append(p)
    spec.personas=people
    task=db.create_task(spec)
    original=AnswerEngine(make_provider(config)).run(task,spec,people[0])
    for p in people:
        r=original.model_copy(deep=True);r.id=uid();r.persona_id=p.id;r.persona_snapshot=p
        db.save_response(r)
    report=quality_report(db,task)
    assert report["summary"]["重複模式紀錄"]==3
    assert all(f["等級"]=="review" for f in report["findings"] if f["代碼"]=="repetition")
    old=db.responses(task)[0]
    new_task=regenerate(db,[old.id],seed=43)
    regenerated_spec = TaskSpec.model_validate_json(db.task(new_task)["spec"])
    assert regenerated_spec.variability == .3 and regenerated_spec.budget_limit == 1.
    assert regenerated_spec.interval_seconds == spec.interval_seconds
    runner=BatchRunner(db,SecretStore());runner.start(new_task);runner.thread.join(timeout=10)
    new=db.responses(new_task)[0]
    assert new.supersedes_id==old.id and new.persona_snapshot==old.persona_snapshot
    assert len(db.responses(task))==3
    assert db.rows("SELECT * FROM quality_actions WHERE response_id=?",(old.id,))


def test_quality_revalidates_bad_format_and_unknown_question(db,saved_response):
    r=saved_response.model_copy(deep=True)
    next(a for a in r.answers if a.question_id=="Q01").value="not-the-age"
    r.answers.append(Answer(question_id="UNKNOWN",value="bad"))
    db.execute("UPDATE responses SET body=? WHERE id=?",(dump(r),r.id))
    report=quality_report(db,r.task_id)
    assert report["summary"]["明確錯誤紀錄"]>=2
    assert {f["代碼"] for f in report["findings"]} >= {"persona","id"}


@pytest.mark.parametrize("ai",[False,True])
def test_motivation_reports_link_original_text_and_safe_exports(db,saved_response,config,ai):
    report=motivation_report(db,saved_response.task_id,"Q12",make_provider(config) if ai else None)
    assert report["dataset"]["dataset_kind"]=="simulation"
    assert report["segments"] and report["themes"]
    assert all(r["response_id"]==saved_response.id for r in report["segments"])
    for row in report["refusal_reasons"]:
        assert row["原文字"] in [r["text"] for r in report["segments"]]
    exports=report_exports(report)
    parsed=json.loads(exports["JSON"])
    assert parsed["method"]==report["method"]
    book=load_workbook(io.BytesIO(exports["Excel"]))
    assert "報告來源與方法" in book.sheetnames and "segments" in book.sheetnames
    assert all(c.data_type != "f" for ws in book for row in ws for c in row)


def test_additive_migration_backup_preserves_v1_data(db,saved_response):
    # Emulate an existing V1 database, then migrate it twice safely.
    with sqlite3.connect(db.path) as conn:
        conn.execute("PRAGMA user_version=0")
        meta=json.loads(conn.execute("SELECT metadata FROM workspace").fetchone()[0]);meta["schema_version"]=1
        conn.execute("UPDATE workspace SET metadata=?",(json.dumps(meta),))
        conn.execute("DROP TABLE schema_migrations")
    before=db.path.parent/"backups"/"before-v2.sqlite3"
    migrate(db.path)
    assert before.exists()
    with sqlite3.connect(before) as conn:
        assert conn.execute("SELECT count(*) FROM responses").fetchone()[0]==1
        assert json.loads(conn.execute("SELECT metadata FROM workspace").fetchone()[0])["schema_version"]==1
    migrate(db.path)
    assert Database(db.path).responses()[0]==saved_response
    assert db.metadata()["schema_version"]==4
    assert len(db.rows("SELECT * FROM schema_migrations"))==3


def test_long_report_metadata_is_preserved_without_excel_truncation():
    from survey_sim.reports import report_tables
    content="x"*70000
    rows=report_tables({"long_snapshot":content})["報告來源與方法"]
    assert len(rows)==3 and ''.join(rows["內容"])==content
    book=load_workbook(io.BytesIO(report_exports({"long_snapshot":content})["Excel"]))
    assert ''.join(row[2].value for row in list(book["報告來源與方法"])[1:])==content
