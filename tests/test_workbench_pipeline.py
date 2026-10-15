import json
import sqlite3
import zipfile
from contextlib import nullcontext

import pytest
from openpyxl import load_workbook

from survey_sim.automation import FormAutomation
from survey_sim.batch import BatchRunner
from survey_sim.engine import AnswerEngine
from survey_sim.errors import AppError, AutomationError
from survey_sim.examples import general_example
from survey_sim.migrations import migrate
from survey_sim.models import Answer, FormTarget, ProviderConfig, TaskSpec
from survey_sim.mock_form import MockFormServer
from survey_sim.personas import PersonaGenerator
from survey_sim.providers import make_provider
from survey_sim.security import SecretStore
from survey_sim.storage import Database, dump
from survey_sim.v2_models import AgeBand, PersonaConditions, WorkbenchSpec
from survey_sim.validator import AnswerValidator
from survey_sim.workbench import WorkbenchService


def execute(service, runner, job_id):
    service.start(job_id, runner)
    runner.thread.join(timeout=30)
    assert not runner.thread.is_alive()
    return service.job(job_id)


@pytest.fixture
def no_browser_server(monkeypatch):
    monkeypatch.setattr("survey_sim.staged_survey.local_form_available", lambda *args: nullcontext())


def test_one_job_generates_controlled_people_answers_and_three_exports(db, config):
    service = WorkbenchService(db, SecretStore())
    runner = BatchRunner(db, service.secrets)
    q = general_example()
    job_id = service.create(WorkbenchSpec(questionnaire=q, provider=config, generate_count=4, interval_seconds=0,
        conditions=PersonaConditions(age_bands=[AgeBand(minimum=30, maximum=30)], occupations={"工程師": 1},
            income_range=(40000, 40000), interests=["閱讀"], genders={"女性": 1, "男性": 1})))
    assert execute(service, runner, job_id)["status"] == "completed"
    report = service.report(job_id)
    assert report["摘要"]["回答數"] == 4
    assert len({p["id"] for p in report["人物"]}) == 4
    assert all(p["background"]["age"] == 30 and p["background"]["occupation"] == "工程師" for p in report["人物"])
    assert sorted(p["background"]["gender"] for p in report["人物"]) == ["女性", "女性", "男性", "男性"]
    assert all(r["dataset_kind"] == "simulation" for r in report["回答紀錄"])
    assert not db.rows("SELECT * FROM submissions")
    folder = db.path.parent / "job-results" / job_id
    assert json.loads((folder / "simulation.json").read_bytes())["dataset"]["dataset_kind"] == "simulation"
    assert "回答明細" in load_workbook(folder / "simulation.xlsx").sheetnames
    with zipfile.ZipFile(folder / "simulation.zip") as archive:
        assert "回答明細.csv" in archive.namelist()
    before = [r.id for r in db.responses()]
    with pytest.raises(AppError, match="不能重新開始"):
        service.start(job_id, runner)
    assert [r.id for r in Database(db.path).responses()] == before


def test_personas_only_and_reuse_across_questionnaires(db, persona, config):
    service = WorkbenchService(db, SecretStore())
    runner = BatchRunner(db, service.secrets)
    person_job = service.create(WorkbenchSpec(provider=config, generate_count=2, mode="personas"))
    assert execute(service, runner, person_job)["status"] == "completed"
    assert len(db.list("persona")) == 2 and not db.responses()
    for i in range(2):
        q = general_example().model_copy(update={"id": f"reuse-{i}"})
        job_id = service.create(WorkbenchSpec(questionnaire=q, provider=config, personas=[persona], interval_seconds=0))
        assert execute(service, runner, job_id)["status"] == "completed"
    assert len(db.responses()) == 2 and len({r.questionnaire_id for r in db.responses()}) == 2
    assert all(r.persona_id == persona.id for r in db.responses())


@pytest.mark.browser
def test_uncertainty_only_can_submit_to_owned_test_form_without_claiming_old_success(db, persona, config):
    q = general_example()
    db.save("questionnaire", q)
    db.save("persona", persona)
    source = TaskSpec(questionnaire=q, personas=[persona], provider=config, interval_seconds=0)
    response = AnswerEngine(make_provider(config)).run(db.create_task(source), source, persona)
    response.answers[0].uncertain = True
    response.answers, response.issues, response.visited_pages, response.status = AnswerValidator().validate(
        q, persona, response.answers)
    assert response.status == "review" and {i.code for i in response.issues} == {"uncertainty"}
    db.save_response(response)
    server = MockFormServer(db, 0).start()
    try:
        target = FormTarget(name="owned", url=server.url(q.id), questionnaire_id=q.id,
            authorized=True, success_selector="#success[data-receipt]")
        service = WorkbenchService(db, SecretStore())
        runner = BatchRunner(db, service.secrets)
        spec = WorkbenchSpec(questionnaire=q, provider=config, response_ids=[response.id], mode="submit",
            form_target=target, form_allowlist=[target.url], submission_confirmed=True, interval_seconds=0)
        first = service.create(spec)
        assert execute(service, runner, first)["status"] == "completed"
        assert service.report(first)["摘要"]["成功畫面或本地收件"] == 1
        assert len(db.rows("SELECT * FROM local_receipts")) == 1
        assert db.responses(response.task_id)[0].status == "review"
        second = service.create(spec)
        assert execute(service, runner, second)["status"] == "completed"
        assert service.report(second)["摘要"]["成功畫面或本地收件"] == 0
        assert service.report(second)["提交紀錄"] == []
        assert len(db.rows("SELECT * FROM local_receipts")) == 1
    finally:
        server.close()


@pytest.mark.browser
def test_missing_success_marker_with_real_receipt_is_waiting_for_verification(db, persona, config):
    q = general_example()
    db.save("questionnaire", q)
    source = TaskSpec(questionnaire=q, personas=[persona], provider=config, interval_seconds=0)
    response = AnswerEngine(make_provider(config)).run(db.create_task(source), source, persona)
    db.save_response(response)
    server = MockFormServer(db, 0).start()
    try:
        target = FormTarget(name="確認文字不符", url=server.url(q.id), questionnaire_id=q.id,
            authorized=True, success_selector="#wrong-success-marker")
        service = WorkbenchService(db, SecretStore())
        runner = BatchRunner(db, service.secrets)
        job_id = service.create(WorkbenchSpec(questionnaire=q, provider=config, response_ids=[response.id],
            mode="submit", form_target=target, form_allowlist=[target.url], submission_confirmed=True,
            interval_seconds=0))
        assert execute(service, runner, job_id)["status"] == "completed_with_errors"
        assert service.report(job_id)["任務明細"][0]["status"] == "unknown"
        assert service.report(job_id)["摘要"]["提交待查"] == 1
        assert service.report(job_id)["摘要"]["成功畫面或本地收件"] == 0
        assert len(db.rows("SELECT * FROM local_receipts")) == 1
    finally:
        server.close()


def test_generation_failure_does_not_block_others_and_retry_keeps_completed(db, config, monkeypatch):
    original = PersonaGenerator.generate_one
    failed = []
    calls = []
    def once(self, seed, *args, **kwargs):
        calls.append(seed)
        if seed == 43 and not failed:
            failed.append(seed)
            raise AppError("temporary persona error")
        return original(self, seed, *args, **kwargs)
    monkeypatch.setattr(PersonaGenerator, "generate_one", once)
    service = WorkbenchService(db, SecretStore())
    runner = BatchRunner(db, service.secrets)
    job_id = service.create(WorkbenchSpec(questionnaire=general_example(), provider=config, generate_count=3, interval_seconds=0))
    assert execute(service, runner, job_id)["status"] == "completed_with_errors"
    first = {r.id for r in db.responses()}
    assert len(first) == 2
    service.retry(job_id)
    assert execute(service, runner, job_id)["status"] == "completed"
    assert len(db.responses()) == 3 and first < {r.id for r in db.responses()}
    assert calls.count(42) == 1 and calls.count(44) == 1 and calls.count(43) == 2


def test_pause_and_recover_generation_checkpoint_without_repeating_completed(db, config, monkeypatch):
    service = WorkbenchService(db, SecretStore())
    runner = BatchRunner(db, service.secrets)
    original = PersonaGenerator.generate_one
    calls = []
    def pause(self, seed, *args, **kwargs):
        calls.append(seed)
        p = original(self, seed, *args, **kwargs)
        if seed == 42:
            service.control(job_id, "paused")
        return p
    monkeypatch.setattr(PersonaGenerator, "generate_one", pause)
    job_id = service.create(WorkbenchSpec(questionnaire=general_example(), provider=config, generate_count=3, interval_seconds=0))
    assert execute(service, runner, job_id)["status"] == "paused"
    assert len(db.list("persona")) == 1
    Database(db.path).recover()
    assert execute(service, runner, job_id)["status"] == "completed"
    assert calls.count(42) == 1 and len(db.responses()) == 3


def test_form_failure_preserves_answer_and_retry_does_not_call_ai_again(db, persona, config, monkeypatch, no_browser_server):
    q = general_example()
    target = FormTarget(name="local", url=f"http://127.0.0.1:9999/form/{q.id}", questionnaire_id=q.id, authorized=True)
    attempts = []
    def form(self, target, response, *args):
        attempts.append(response.id)
        if len(attempts) == 1:
            raise AutomationError("mapping failed")
        return {"status": "dry_run", "id": "test"}
    monkeypatch.setattr(FormAutomation, "run", form)
    service = WorkbenchService(db, SecretStore())
    runner = BatchRunner(db, service.secrets)
    job_id = service.create(WorkbenchSpec(questionnaire=q, provider=config, personas=[persona], mode="dry_run",
        form_target=target, form_allowlist=[target.url], interval_seconds=0))
    assert execute(service, runner, job_id)["status"] == "completed_with_errors"
    saved = db.responses()[0]
    monkeypatch.setattr(AnswerEngine, "run", lambda *a, **k: pytest.fail("completed answers must not be regenerated"))
    monkeypatch.setattr("survey_sim.batch.make_provider", lambda *a, **k: pytest.fail("saved answers require no AI client"))
    service.retry(job_id)
    assert execute(service, runner, job_id)["status"] == "completed"
    assert attempts == [saved.id, saved.id] and len(db.responses()) == 1


@pytest.mark.parametrize("dry_run", [True, False])
@pytest.mark.browser
def test_full_job_real_browser_local_form_and_existing_answer_batch(db, config, questionnaire, dry_run):
    q = questionnaire.model_copy(deep=True)
    q.consistency_rules = []
    for question in q.questions:
        question.trait = ""
    db.save("questionnaire", q)
    server = MockFormServer(db, 0).start()
    try:
        target = FormTarget(name="owned local", url=server.url(q.id), questionnaire_id=q.id, authorized=True)
        service = WorkbenchService(db, SecretStore())
        runner = BatchRunner(db, service.secrets)
        job_id = service.create(WorkbenchSpec(questionnaire=q, provider=config, generate_count=3, interval_seconds=0,
            mode="dry_run" if dry_run else "submit", form_target=target, form_allowlist=[target.url], submission_confirmed=not dry_run))
        job = execute(service, runner, job_id)
        assert job["status"] == "completed", job["error"]
        assert len(db.responses()) == 3
        assert len(db.rows("SELECT * FROM local_receipts")) == (0 if dry_run else 3)
        assert len(db.rows("SELECT * FROM submissions")) == 3
        if not dry_run:
            again = service.create(WorkbenchSpec(provider=config, response_ids=[r.id for r in db.responses()],
                mode="submit", form_target=target, form_allowlist=[target.url], submission_confirmed=True, interval_seconds=0))
            assert execute(service, runner, again)["status"] == "completed"
            assert all(i["status"] == "skipped" for i in service.report(again)["任務明細"])
            assert len(db.rows("SELECT * FROM local_receipts")) == 3
            assert len(db.rows("SELECT * FROM submissions")) == 3
    finally:
        server.close()


def test_unauthorized_reference_and_unconfirmed_submissions_are_blocked_before_queue(db, persona, config):
    service = WorkbenchService(db, SecretStore())
    q = general_example()
    target = FormTarget(name="test", url="https://example.test/form", questionnaire_id=q.id, kind="custom", success_text="received")
    with pytest.raises(ValueError, match="明確確認"):
        WorkbenchSpec(questionnaire=q, provider=config, personas=[persona], mode="submit", form_target=target, submission_confirmed=True)
    target.authorized = True
    q.purpose = "reference"
    spec = WorkbenchSpec(questionnaire=q, provider=config, personas=[persona], mode="submit", form_target=target,
        form_allowlist=[target.url], submission_confirmed=True)
    with pytest.raises(AppError, match="正式研究"):
        service.create(spec)
    assert not db.rows("SELECT * FROM workbench_jobs") and not db.rows("SELECT * FROM submissions")


def test_uncertainty_is_review_and_never_silently_treated_as_verified(questionnaire, persona, config):
    response = AnswerEngine(make_provider(config)).run("uncertain", TaskSpec(questionnaire=questionnaire, personas=[persona], provider=config), persona)
    response.answers[0].uncertain = True
    _, issues, _, status = AnswerValidator().validate(questionnaire, persona, response.answers)
    assert status == "review" and any(i.code == "uncertainty" for i in issues)


def test_v3_migration_keeps_entities_answers_and_provider_configs(db, saved_response, config):
    db.save("provider", config)
    entities = db.rows("SELECT * FROM entities ORDER BY kind,id")
    before = db.responses()[0].model_dump_json()
    with sqlite3.connect(db.path) as conn:
        conn.execute("DROP TABLE workbench_jobs")
        conn.execute("DELETE FROM schema_migrations WHERE version=4")
        conn.execute("PRAGMA user_version=3")
    migrate(db.path)
    migrate(db.path)
    assert Database(db.path).responses()[0].model_dump_json() == before
    assert (db.path.parent / "backups/before-v4.sqlite3").exists()
    assert db.metadata()["schema_version"] == 4
    assert db.rows("SELECT name FROM sqlite_master WHERE name='workbench_jobs'")
    assert db.rows("SELECT * FROM entities ORDER BY kind,id") == entities
    assert db.get("provider", config.id) == config


def test_cancel_stops_remaining_generation_and_keeps_finished_persona(db, config, monkeypatch):
    service = WorkbenchService(db, SecretStore())
    runner = BatchRunner(db, service.secrets)
    original = PersonaGenerator.generate_one
    def cancel(self, seed, *args, **kwargs):
        person = original(self, seed, *args, **kwargs)
        service.control(job_id, "cancelled")
        return person
    monkeypatch.setattr(PersonaGenerator, "generate_one", cancel)
    job_id = service.create(WorkbenchSpec(provider=config, generate_count=3, mode="personas"))
    assert execute(service, runner, job_id)["status"] == "cancelled"
    assert len(db.list("persona")) == 1
    assert [i["status"] for i in json.loads(service.job(job_id)["checkpoint"])["generation"]] == ["completed", "cancelled", "cancelled"]


def test_transient_load_retry_is_bounded_and_does_not_regenerate(db, persona, config, monkeypatch, no_browser_server):
    config.retries = 1
    q = general_example()
    target = FormTarget(name="owned", url=f"http://127.0.0.1:9999/form/{q.id}", questionnaire_id=q.id, authorized=True)
    calls = []
    def fail(self, target, response, *args):
        calls.append(response.id)
        raise AutomationError("temporary load timeout", retryable=True)
    monkeypatch.setattr(FormAutomation, "run", fail)
    service = WorkbenchService(db, SecretStore())
    runner = BatchRunner(db, service.secrets)
    job_id = service.create(WorkbenchSpec(questionnaire=q, provider=config, personas=[persona], mode="dry_run",
        form_target=target, form_allowlist=[target.url], interval_seconds=0))
    assert execute(service, runner, job_id)["status"] == "completed_with_errors"
    assert len(calls) == 2 and len(set(calls)) == 1 and len(db.responses()) == 1


def test_unknown_submission_blocks_retry_without_another_send(db, persona, config, monkeypatch, no_browser_server):
    q = general_example()
    for question in q.questions:
        question.trait = ""
    target = FormTarget(name="owned", url=f"http://127.0.0.1:9999/form/{q.id}", questionnaire_id=q.id, authorized=True)
    sends = []
    def unknown(self, target, response, allowlist, dry_run, *args):
        token = db.reserve_submission(response.id, target, False)
        sends.append(response.id)
        db.finish_submission(token, "unknown", "test acknowledgement lost")
        raise AutomationError("提交待查，禁止重送")
    monkeypatch.setattr(FormAutomation, "run", unknown)
    service = WorkbenchService(db, SecretStore())
    runner = BatchRunner(db, service.secrets)
    job_id = service.create(WorkbenchSpec(questionnaire=q, provider=config, personas=[persona], mode="submit",
        form_target=target, form_allowlist=[target.url], submission_confirmed=True, interval_seconds=0))
    assert execute(service, runner, job_id)["status"] == "completed_with_errors"
    service.retry(job_id)
    assert execute(service, runner, job_id)["status"] == "completed_with_errors"
    assert len(sends) == 1 and len(db.responses()) == 1
    assert db.rows("SELECT status FROM submissions")[0]["status"] == "unknown"


def test_all_generation_failures_still_produce_an_automatic_failure_report(db, config, monkeypatch):
    monkeypatch.setattr(PersonaGenerator, "generate_one", lambda *a, **k: (_ for _ in ()).throw(AppError("API unavailable")))
    service = WorkbenchService(db, SecretStore())
    runner = BatchRunner(db, service.secrets)
    job_id = service.create(WorkbenchSpec(questionnaire=general_example(), provider=config, generate_count=2))
    assert execute(service, runner, job_id)["status"] == "failed"
    report = json.loads((db.path.parent / "job-results" / job_id / "simulation.json").read_bytes())
    assert report["status"] == "failed" and len(report["人物生成紀錄"]) == 2
    assert all(i["error"] == "API unavailable" for i in report["人物生成紀錄"])
    assert report["摘要"]["失敗"] == 2
