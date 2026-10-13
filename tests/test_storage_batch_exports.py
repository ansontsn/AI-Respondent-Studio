import io
import json
import threading
import time
import zipfile

import pandas as pd
import pytest
from openpyxl import load_workbook

from survey_sim.analysis import export_csv_zip, export_excel, export_json, concentration
from survey_sim.batch import BatchRunner
from survey_sim.errors import AppError
from survey_sim.models import TaskSpec
from survey_sim.mock_ai import MockProvider
from survey_sim.security import SecretStore
from survey_sim.storage import Database


def test_restart_reads_history_and_snapshots(db, saved_response):
    p = saved_response.persona_snapshot.model_copy(deep=True)
    p.name = "已修改人物"
    db.save("persona", p)
    reopened = Database(db.path)
    assert reopened.responses()[0] == saved_response
    assert reopened.get("persona", p.id).name != reopened.responses()[0].persona_snapshot.name
    assert reopened.metadata()["dataset_kind"] == "simulation"
    assert len(reopened.rows("SELECT * FROM answers")) == len(saved_response.answers)


def test_exports_cross_reference_source_and_safe_excel(db, saved_response):
    saved_response.persona_snapshot.name = "=HYPERLINK(\"bad\")"
    db.save("persona", saved_response.persona_snapshot)
    json_data = json.loads(export_json(db, saved_response.task_id))
    assert json_data["dataset"]["dataset_kind"] == "simulation"
    assert json_data["responses"][0]["persona_id"] == saved_response.persona_id
    book = load_workbook(io.BytesIO(export_excel(db, saved_response.task_id)))
    assert {"資料集來源", "人物", "回答", "回答明細", "驗證報告", "執行紀錄", "題目對照"} <= set(book.sheetnames)
    assert book["回答"].freeze_panes == "A2"
    assert all(c.data_type != "f" for ws in book for row in ws for c in row)
    with zipfile.ZipFile(io.BytesIO(export_csv_zip(db))) as z:
        assert json.loads(z.read("dataset_metadata.json"))["source"]
        people = pd.read_csv(io.BytesIO(z.read("人物.csv")))
        assert people["name"].iloc[0].startswith("'=HYPERLINK")
        answers = pd.read_csv(io.BytesIO(z.read("回答.csv")))
        assert answers["persona_id"].iloc[0] == saved_response.persona_id
        assert "Q01" in answers
    assert concentration([saved_response])["不同答案組合"] == 1


def test_database_backup(db, saved_response, tmp_path):
    path = tmp_path / "backup.sqlite3"
    db.backup(path)
    assert Database(path).responses() == db.responses()


def test_export_keeps_deleted_historical_persona_and_task_scope(db, saved_response):
    db.delete("persona", saved_response.persona_id)
    from survey_sim.analysis import tables
    exported = tables(db)
    assert saved_response.persona_id in list(exported["人物"]["persona_id"])


def test_unconfirmed_questionnaire_rejected(db, spec):
    spec.questionnaire.confirmed = False
    with pytest.raises(AppError): db.create_task(spec)


def test_restart_recovery_preserves_checkpoint_and_unknown_submission(db, spec, saved_response):
    task = db.create_task(spec)
    db.task_status(task, "running")
    item = db.claim_item(task)
    db.checkpoint(task, item["persona_id"], {"answers": [], "next_page": "P2"})
    from survey_sim.models import FormTarget
    target = FormTarget(name="test", url="http://127.0.0.1:8765/form/test", questionnaire_id=spec.questionnaire.id)
    id = db.reserve_submission(saved_response.id, target, False)
    db.recover()
    assert db.task(task)["status"] == "paused"
    assert json.loads(db.rows("SELECT checkpoint FROM task_items WHERE task_id=?", (task,))[0]["checkpoint"])["next_page"] == "P2"
    assert db.rows("SELECT status FROM submissions WHERE id=?", (id,))[0]["status"] == "unknown"


def test_batch_failure_retains_completed_and_retry_only_failed(db, spec, config, monkeypatch):
    from survey_sim.personas import PersonaGenerator
    from survey_sim.providers import make_provider
    people = PersonaGenerator(make_provider(config)).generate_batch(3, 110, "一般", {"AI 自動混合": 1})
    spec.personas = people
    task = db.create_task(spec)
    class FailOne(MockProvider):
        def generate(self, operation, payload, system):
            if operation == "answer_page" and payload["persona"]["id"] == people[1].id:
                raise RuntimeError("simulated outage")
            return super().generate(operation, payload, system)
    monkeypatch.setattr("survey_sim.batch.make_provider", lambda *a: FailOne(config))
    runner = BatchRunner(db, SecretStore())
    runner.start(task)
    runner.thread.join(timeout=10)
    assert len(db.responses(task)) == 2
    ids = [r.id for r in db.responses(task)]
    assert db.task(task)["status"] == "completed_with_errors"
    monkeypatch.setattr("survey_sim.batch.make_provider", lambda *a: MockProvider(config))
    db.retry_failed(task)
    runner.start(task)
    runner.thread.join(timeout=10)
    assert len(db.responses(task)) == 3
    assert set(ids) <= {r.id for r in db.responses(task)}
    assert db.task(task)["status"] == "completed"


def test_pause_stop_and_resume_at_page_boundary(db, spec, config, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    class Slow(MockProvider):
        calls = 0
        def generate(self, operation, payload, system):
            if operation == "answer_page" and self.calls == 0:
                self.calls += 1
                entered.set()
                release.wait(timeout=5)
            return super().generate(operation, payload, system)
    monkeypatch.setattr("survey_sim.batch.make_provider", lambda *a: Slow(config))
    runner = BatchRunner(db, SecretStore())
    task = db.create_task(spec)
    runner.start(task)
    assert entered.wait(timeout=5)
    runner.pause(task)
    release.set()
    runner.thread.join(timeout=5)
    assert db.task(task)["status"] == "paused"
    assert json.loads(db.rows("SELECT checkpoint FROM task_items WHERE task_id=?", (task,))[0]["checkpoint"])["answers"]
    runner.stop(task)
    assert db.task(task)["status"] == "stopped"
    runner.start(task)
    runner.thread.join(timeout=5)
    assert db.task(task)["status"] == "completed"
