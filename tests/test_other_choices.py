import io
import json
import sqlite3
import zipfile

import httpx
import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from survey_sim.analysis import display_answer, distribution, export_csv_zip, export_excel, export_json
from survey_sim.automation import FormAutomation, launch_browser
from survey_sim.engine import AnswerEngine
from survey_sim.google_import import parse_official_form, parse_public_metadata, stable_option
from survey_sim.mock_ai import MockProvider
from survey_sim.mock_form import MockFormServer
from survey_sim.models import Answer, FormTarget, Option, Page, Question, Questionnaire, TaskSpec
from survey_sim.parsers import parse_google_dom
from survey_sim.storage import Database
from survey_sim.validator import AnswerValidator, value_errors
from survey_sim.website_experience import compact_payload


def other_question(kind="multiple", required=True):
    return Question(id="Q_other", title="收藏類型", type=kind, required=required,
                    options=[Option(id="kpop", label="K-POP 偶像小卡"), Option(id="other", label="其他", is_other=True)])


@pytest.mark.parametrize("kind,expected", [(2, "single"), (4, "multiple")])
@pytest.mark.parametrize("other_label", ["", "其他"])
def test_public_metadata_recognizes_empty_other_label(kind, expected, other_label):
    form = [None] * 9
    form[1] = [[1654015254, "收藏類型", None, kind,
                [[123, [["K-POP 偶像小卡"], [other_label, None, None, None, 1]], 1]]]]
    form[8] = "測試匯入"
    doc = [None, form]
    html = "<script>FB_PUBLIC_LOAD_DATA_=" + json.dumps(doc, ensure_ascii=False) + ";</script>"
    q = parse_public_metadata(html, "https://docs.google.com/forms/d/e/test/viewform")
    item = q.questions[0]
    assert item.type == expected and item.id == "G_1654015254"
    assert item.options[0].id == stable_option("K-POP 偶像小卡")
    assert item.options[1].label == "其他" and item.options[1].is_other
    assert not any("尚未支援" in warning or "無法辨識" in warning for warning in q.import_warnings)
    assert not q.confirmed


@pytest.mark.parametrize("kind,expected", [("RADIO", "single"), ("CHECKBOX", "multiple")])
def test_official_json_preserves_other_and_roundtrip(kind, expected):
    data = {"info": {"title": "收藏類型"}, "items": [{"itemId": "i1", "questionItem": {"question": {
        "questionId": "q1", "choiceQuestion": {"type": kind, "options": [{"value": "一般"}, {"isOther": True}]}}}}]}
    q = parse_official_form(data)
    assert q.questions[0].type == expected
    assert q.questions[0].options[-1].is_other
    assert Questionnaire.model_validate_json(q.model_dump_json()) == q


def test_dom_uses_explicit_other_marker_and_does_not_guess_from_label():
    html = '<div role="listitem"><div role="heading">收藏類型</div><div role="checkbox" data-value="K-POP"></div><div role="checkbox" data-value="__other_option__" aria-label="其他："></div><input type="text"></div>'
    q = parse_google_dom(html, "url").questions[0]
    assert q.type == "multiple" and q.options[-1].is_other and q.options[-1].label == "其他"
    ordinary = html.replace('data-value="__other_option__"', 'data-value="其他"').replace('<input type="text">', '')
    assert not any(o.is_other for o in parse_google_dom(ordinary, "url").questions[0].options)
    unknown = html.replace('data-value="__other_option__"', 'data-value="其他"')
    assert parse_google_dom(unknown, "url").questions[0].type == "unsupported"


@pytest.mark.parametrize("value,text,valid", [
    (["kpop", "other"], "棒球球員卡", True), (["other"], "  棒球卡  ", True),
    (["kpop"], "", True), (["other"], "", False), (["other"], "   ", False),
    (["kpop"], "棒球卡", False), (["other", "other"], "棒球卡", False),
    (["unknown", "other"], "棒球卡", False), (None, "棒球卡", False),
    (["other"], "字" * 2001, False), (["other"], 123, False),
])
def test_other_text_validation(value, text, valid):
    assert (not value_errors(other_question(), value, text)) == valid


def test_limits_single_and_optional_other():
    q = other_question()
    q.max_choices = 1
    assert value_errors(q, ["kpop", "other"], "棒球卡")
    assert not value_errors(other_question("single"), "other", "棒球卡")
    assert not value_errors(other_question(required=False), [], "")
    assert value_errors(other_question(required=False), ["other"], "")
    assert value_errors(Question(id="text", title="文字", type="short_text"), "一般文字", "不應附帶")


def build_response(db, persona, config):
    class FixedOther(MockProvider):
        def generate(self, operation, payload, system):
            assert "other_text" in system
            result = super().generate(operation, payload, system)
            for a in result.data["answers"]:
                a.update(value=["kpop", "other"] if a["question_id"] == "Q_other" else "other",
                         other_text="棒球球員卡", rationale="測試其他文字保存。")
            return result
    single = other_question("single").model_copy(update={"id": "Q_single"})
    q = Questionnaire(title="其他文字測試", confirmed=True, pages=[Page(id="P1", title="收藏", questions=[other_question(), single])])
    spec = TaskSpec(questionnaire=q, personas=[persona], provider=config, interval_seconds=0)
    db.save("questionnaire", q)
    task_id = db.create_task(spec)
    response = AnswerEngine(FixedOther(config)).run(task_id, spec, persona)
    db.save_response(response)
    return response


def test_other_roundtrip_storage_exports_and_counts(db, persona, config):
    response = build_response(db, persona, config)
    loaded = Database(db.path).responses(response.task_id)[0]
    assert loaded == response and all(a.other_text == "棒球球員卡" for a in loaded.answers)
    assert all(a["other_text"] == "棒球球員卡" for a in db.rows("SELECT * FROM answers"))
    assert display_answer(response.questionnaire_snapshot.questions[0], ["kpop", "other"], "棒球球員卡") == "K-POP 偶像小卡、其他：棒球球員卡"
    counts = distribution([response], "Q_other").set_index("選項")["回答次數"].to_dict()
    assert counts == {"K-POP 偶像小卡": 1, "其他": 1}
    data = json.loads(export_json(db))
    assert data["responses"][0]["answers"][0]["other_text"] == "棒球球員卡"
    assert data["dataset"]["dataset_kind"] == "simulation"
    with zipfile.ZipFile(io.BytesIO(export_csv_zip(db))) as archive:
        wide = pd.read_csv(io.BytesIO(archive.read("回答.csv")))
        detail = pd.read_csv(io.BytesIO(archive.read("回答明細.csv")))
        assert wide.iloc[0]["Q_other__other_text"] == "棒球球員卡"
        assert detail.iloc[0]["other_text"] == "棒球球員卡"
        assert "其他：棒球球員卡" in detail.iloc[0]["answer_text"]
    book = load_workbook(io.BytesIO(export_excel(db)))
    headers = [cell.value for cell in book["回答"][1]]
    assert book["回答"].cell(2, headers.index("Q_other__other_text") + 1).value == "棒球球員卡"
    # Existing stored answers remain readable without the new optional field.
    assert Answer.model_validate({"question_id": "old", "value": ["kpop"]}).other_text == ""


def test_engine_repairs_missing_other_text_and_preserves_checkpoint(persona, config):
    class Repair(MockProvider):
        calls = 0
        def generate(self, operation, payload, system):
            self.calls += 1
            if self.calls == 2:
                assert "補充文字" in payload["validation_feedback"]
            result = super().generate(operation, payload, system)
            result.data["answers"][0].update(value=["other"], other_text="" if self.calls == 1 else "棒球卡")
            return result
    q = Questionnaire(title="測試", pages=[Page(id="P1", title="收藏", questions=[other_question()])])
    spec = TaskSpec(questionnaire=q, personas=[persona], provider=config)
    checkpoints = []
    r = AnswerEngine(Repair(config)).run("test", spec, persona, on_checkpoint=checkpoints.append)
    assert r.answers[0].valid and r.answers[0].other_text == "棒球卡"
    assert checkpoints[-1]["answers"][0]["other_text"] == "棒球卡"
    payload = {"questionnaire": q.model_dump(), "page": q.pages[0].model_dump(), "previous_answers": [r.answers[0].model_dump()]}
    assert compact_payload(payload)["previous_answers"][0]["other_text"] == "棒球卡"
    r.answers[0].other_text = ""
    assert AnswerValidator().validate(q, persona, r.answers)[-1] == "error"


def test_migration_from_v2_adds_text_and_preserves_backup(db, saved_response):
    with sqlite3.connect(db.path) as conn:
        conn.execute("ALTER TABLE answers DROP COLUMN other_text")
        conn.execute("CREATE TABLE local_receipts(id TEXT PRIMARY KEY,questionnaire_id TEXT,token TEXT,answers TEXT,created_at TEXT)")
        conn.execute("INSERT INTO local_receipts VALUES('old','q','token','{}','date')")
        conn.execute("DELETE FROM schema_migrations WHERE version=3")
        conn.execute("PRAGMA user_version=2")
    upgraded = Database(db.path)
    assert upgraded.responses()[0] == saved_response
    assert all(row["other_text"] == "" for row in upgraded.rows("SELECT * FROM answers"))
    assert upgraded.rows("SELECT other_text FROM local_receipts")[0]["other_text"] == "{}"
    assert (db.path.parent / "backups" / "before-v3.sqlite3").exists()
    assert len(Database(db.path).rows("SELECT * FROM schema_migrations")) == 3


def test_server_rejects_missing_or_unselected_other_text(db, persona, config):
    r = build_response(db, persona, config)
    server = MockFormServer(db, 0).start()
    try:
        url = server.url(r.questionnaire_id).replace("/form/", "/submit/")
        payload = {"token": "other-test", "answers": {a.question_id: a.value for a in r.answers}, "visited": ["P1"]}
        assert httpx.post(url, json=payload).status_code == 400
        payload["other_text"] = {a.question_id: a.other_text for a in r.answers}
        payload["answers"]["Q_other"] = ["kpop"]
        assert httpx.post(url, json=payload).status_code == 400
        payload["answers"]["Q_other"] = ["kpop", "other"]
        assert httpx.post(url, json=payload).status_code == 200
        stored = db.rows("SELECT * FROM local_receipts")[0]
        assert json.loads(stored["other_text"])["Q_other"] == "棒球球員卡"
    finally:
        server.close()


def test_preview_shows_other_field():
    q = Questionnaire(title="測試", pages=[Page(id="P1", title="收藏", questions=[other_question()])])
    code = "from survey_sim.models import Questionnaire\nfrom survey_sim.questionnaire_review import preview\npreview(Questionnaire.model_validate_json(" + repr(q.model_dump_json()) + "))"
    app = AppTest.from_string(code).run()
    assert not app.exception
    assert app.multiselect[0].options == ["K-POP 偶像小卡", "其他"]
    assert any("勾選後填寫" in field.label for field in app.text_input)


@pytest.mark.browser
def test_browser_local_preserves_selected_and_text(db, persona, config, tmp_path):
    r = build_response(db, persona, config)
    server = MockFormServer(db, 0).start()
    try:
        target = FormTarget(name="其他文字", url=server.url(r.questionnaire_id), questionnaire_id=r.questionnaire_id, authorized=True)
        automation = FormAutomation(db, tmp_path)
        assert automation.run(target, r, [target.url], True)["status"] == "dry_run"
        assert not db.rows("SELECT * FROM local_receipts")
        assert automation.run(target, r, [target.url], False)["status"] == "success"
        receipt = db.rows("SELECT * FROM local_receipts")[0]
        assert json.loads(receipt["answers"]) == {a.question_id: a.value for a in r.answers}
        assert json.loads(receipt["other_text"]) == {a.question_id: a.other_text for a in r.answers}
    finally:
        server.close()


@pytest.mark.browser
def test_google_style_other_control_with_local_dom(db, tmp_path):
    from playwright.sync_api import sync_playwright
    q = other_question()
    html = '<div role="listitem"><div role="heading">收藏類型</div><div role="checkbox" aria-label="K-POP 偶像小卡" aria-checked="false" onclick="this.setAttribute(\'aria-checked\',\'true\')">K-POP 偶像小卡</div><div role="checkbox" data-value="__other_option__" aria-label="其他：" aria-checked="false" onclick="this.setAttribute(\'aria-checked\',\'true\')">其他：</div><input type="text" aria-label="其他回答"></div>'
    with sync_playwright() as pw:
        browser = launch_browser(pw)
        try:
            page = browser.new_page()
            page.set_content(html)
            FormAutomation(db, tmp_path)._fill(page, q, ["kpop", "other"], "google", "棒球球員卡")
            assert page.locator('[role="checkbox"][aria-checked="true"]').count() == 2
            assert page.locator('input[type="text"]').input_value() == "棒球球員卡"
        finally:
            browser.close()
